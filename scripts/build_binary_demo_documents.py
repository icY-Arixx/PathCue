from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors


ROOT = Path(__file__).resolve().parents[1]
KNOWLEDGE = ROOT / "data" / "knowledge_base"


def set_run_font(run, name: str = "Microsoft YaHei", size: float = 11, bold: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def build_docx() -> Path:
    output = KNOWLEDGE / "traffic_notice_lifecycle.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(0.85)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("城市交通事件公告数据治理规范"), size=20, bold=True)
    intro = doc.add_paragraph()
    set_run_font(intro.add_run("本规范用于统一道路施工、突发事件和临时交通管制公告的采集、审核、发布、更新与归档流程，确保知识库中的时间、位置和影响范围可追溯。"))

    sections = [
        ("一 数据字段", [
            "每条公告必须包含事件编号、公告标题、道路名称、行政区、起止坐标、影响方向、开始时间、预计结束时间、事件类型、影响等级、绕行建议、来源单位、发布时间和更新时间。",
            "坐标统一采用 GCJ-02，经纬度保留六位小数；道路名称应关联标准路网 ID，行政区使用国家统计区划代码。",
        ]),
        ("二 发布与审核", [
            "采集任务首先校验必填字段、时间逻辑和空间范围，再由审核人员确认来源可靠性。审核通过后生成版本号并发布到检索索引。",
            "对缺失结束时间、坐标落在城市边界外、开始时间晚于结束时间等异常记录，应进入人工复核队列，不得直接发布。",
        ]),
        ("三 增量更新", [
            "系统以事件编号作为业务主键，以内容哈希和更新时间识别增量。字段未变化时跳过向量化；正文、时间或空间范围变化时删除旧分块并写入新版本。",
            "公告撤销后保留历史版本和撤销原因，同时从在线检索集合中移除有效分块。每次更新记录操作人、更新时间、变更字段和前后版本号。",
        ]),
        ("四 质量指标", [
            "每日输出元数据完整率、坐标有效率、时间逻辑通过率、重复公告率、更新及时率和人工复核率。重复公告通过事件编号、标题相似度与空间时间重叠联合识别。",
            "当元数据完整率低于 98% 或坐标有效率低于 99% 时触发质量告警，并暂停问题来源的自动入库。",
        ]),
    ]
    for heading, paragraphs in sections:
        p = doc.add_paragraph(style="Heading 1")
        set_run_font(p.add_run(heading), size=14, bold=True)
        for text in paragraphs:
            body = doc.add_paragraph(style="List Bullet")
            set_run_font(body.add_run(text))
    doc.core_properties.title = "城市交通事件公告数据治理规范"
    doc.core_properties.author = "Aris Xu"
    doc.save(output)
    return output


def build_pdf() -> Path:
    output = KNOWLEDGE / "emergency_accessibility_guide.pdf"
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    font_name = "MicrosoftYaHei"
    pdfmetrics.registerFont(TTFont(font_name, str(font_path), subfontIndex=0))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ZhTitle", parent=styles["Title"], fontName=font_name, fontSize=18, leading=24, alignment=TA_CENTER, textColor=colors.black)
    heading_style = ParagraphStyle("ZhHeading", parent=styles["Heading2"], fontName=font_name, fontSize=13, leading=18, spaceBefore=10, spaceAfter=6, textColor=colors.black)
    body_style = ParagraphStyle("ZhBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=17, spaceAfter=7, textColor=colors.black)
    story = [
        Paragraph("公共交通应急接驳与无障碍出行指引", title_style),
        Spacer(1, 0.18 * inch),
        Paragraph("本指引用于地铁临时停运、极端天气和大型活动客流管控场景，为应急接驳、无障碍服务和乘客信息发布提供统一规则。", body_style),
    ]
    content = [
        ("应急接驳启动", "当线路预计中断超过 20 分钟，运营单位应评估客流并启动接驳预案。接驳点优先设置在停运区段两端车站的公交接驳区，避免占用消防通道和无障碍坡道。"),
        ("信息发布", "公告至少包含受影响线路、停运区段、开始时间、预计恢复时间、接驳点名称、步行距离、发车间隔、末班接驳时间、无障碍车辆数量和现场联系人。信息发生变化后 10 分钟内完成更新。"),
        ("无障碍服务", "轮椅乘客应优先引导至具备低地板车辆和无障碍升降设备的接驳点。视障乘客需要语音播报、连续盲道和工作人员引导；听障乘客需要电子屏和文字提示。"),
        ("空间数据要求", "接驳点坐标使用 GCJ-02，并关联车站出入口 ID、道路可达方向和最近的无障碍通道。步行路径不得穿越封闭施工区域，坡度、台阶和电梯状态需要作为路径属性维护。"),
        ("评估与复盘", "事件结束后统计平均等待时间、接驳运力利用率、无障碍请求响应率、信息更新及时率和乘客投诉量。应将漏报字段、错误坐标和过期公告纳入数据质量问题单。"),
    ]
    for heading, text in content:
        story.append(Paragraph(heading, heading_style))
        story.append(Paragraph(text, body_style))
    data = [
        ["指标", "目标"],
        ["信息更新及时率", "不低于 95%"],
        ["无障碍请求响应率", "不低于 98%"],
        ["接驳点坐标有效率", "不低于 99%"],
    ]
    table = Table(data, colWidths=[2.6 * inch, 2.6 * inch])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([Spacer(1, 0.12 * inch), table])
    SimpleDocTemplate(str(output), pagesize=letter, leftMargin=0.85 * inch, rightMargin=0.85 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch, title="公共交通应急接驳与无障碍出行指引", author="Aris Xu").build(story)
    return output


def write_sidecar(path: Path, payload: dict[str, str]) -> None:
    path.with_suffix(path.suffix + ".meta.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    KNOWLEDGE.mkdir(parents=True, exist_ok=True)
    docx = build_docx()
    pdf = build_pdf()
    write_sidecar(docx, {"doc_id": "traffic_notice_lifecycle", "title": "城市交通事件公告数据治理规范", "city": "武汉", "category": "数据治理", "updated_at": "2026-09-11", "authority": "demo"})
    write_sidecar(pdf, {"doc_id": "emergency_accessibility", "title": "公共交通应急接驳与无障碍出行指引", "city": "武汉", "category": "公共交通", "updated_at": "2026-09-11", "authority": "demo"})
    print(docx)
    print(pdf)
