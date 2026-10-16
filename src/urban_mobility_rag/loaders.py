from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Iterable

from .models import SourceDocument
from .text import normalize_text

SUPPORTED_EXTENSIONS = {".md", ".txt", ".json", ".jsonl", ".csv", ".docx", ".pdf"}


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _parse_front_matter(text: str) -> tuple[dict[str, str], str]:
    if not text.startswith("---\n"):
        return {}, text
    _, header, body = text.split("---", 2)
    metadata: dict[str, str] = {}
    for line in header.splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            metadata[key.strip()] = value.strip().strip('"')
    return metadata, body.strip()


def _load_text(path: Path) -> tuple[dict[str, str], str]:
    raw = path.read_text(encoding="utf-8-sig")
    return _parse_front_matter(raw)


def _load_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    return "\n".join("；".join(f"{key}：{value}" for key, value in row.items()) for row in rows)


def _load_json(path: Path) -> str:
    if path.suffix.lower() == ".jsonl":
        records = [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
    else:
        records = json.loads(path.read_text(encoding="utf-8-sig"))
    return json.dumps(records, ensure_ascii=False, indent=2)


def _load_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Reading DOCX requires the optional 'documents' dependencies") from exc
    document = Document(path)
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        blocks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(blocks)


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Reading PDF requires the optional 'documents' dependencies") from exc
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _load_sidecar_metadata(path: Path) -> dict[str, str]:
    sidecar = path.with_suffix(path.suffix + ".meta.json")
    if not sidecar.exists():
        return {}
    payload = json.loads(sidecar.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, dict):
        raise ValueError(f"Metadata sidecar must be a JSON object: {sidecar}")
    return payload


def load_document(path: Path) -> SourceDocument:
    suffix = path.suffix.lower()
    metadata: dict[str, str] = {}
    if suffix in {".md", ".txt"}:
        metadata, text = _load_text(path)
    elif suffix in {".json", ".jsonl"}:
        text = _load_json(path)
    elif suffix == ".csv":
        text = _load_csv(path)
    elif suffix == ".docx":
        text = _load_docx(path)
    elif suffix == ".pdf":
        text = _load_pdf(path)
    else:
        raise ValueError(f"Unsupported document type: {suffix}")
    metadata.update(_load_sidecar_metadata(path))
    doc_id = metadata.get("doc_id") or path.stem
    title = metadata.get("title") or path.stem.replace("_", " ")
    metadata.update({"sha256": file_digest(path), "path": str(path)})
    return SourceDocument(doc_id, title, normalize_text(text), str(path), suffix.lstrip("."), metadata)


def iter_documents(root: Path) -> Iterable[SourceDocument]:
    for path in sorted(root.rglob("*")):
        if path.is_file() and not path.name.endswith(".meta.json") and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield load_document(path)
